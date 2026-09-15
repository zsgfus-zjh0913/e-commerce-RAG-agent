import os
import sys
import json
import re
import hashlib

_rag_libs = os.environ.get("RAG_LIBS_PATH", r"C:\rag_libs")
if os.path.isdir(_rag_libs):
    import site
    site.addsitedir(_rag_libs)
    if _rag_libs not in sys.path:
        sys.path.append(_rag_libs)

from pathlib import Path
import numpy as np

os.environ.setdefault("HF_ENDPOINT", "https://hf-mirror.com")

import jieba

CUSTOM_WORDS = [
    "尺码", "型号", "面料", "材质", "商品编号", "库存",
    "包邮", "退换货", "售后服务", "发货时间", "产地",
    "男装", "女装", "童装", "鞋类", "配饰", "电子产品",
    "家居用品", "食品", "美妆", "运动户外",
    "蓝牙耳机", "无线耳机", "声科达", "SC-505",
    "T恤", "卫衣", "冲锋衣", "连衣裙", "牛仔裤",
    "胸围", "腰围", "臀围", "肩宽", "衣长", "袖长",
    "退款", "换货", "退货", "保修", "质保",
]

for w in CUSTOM_WORDS:
    jieba.add_word(w)

SUPPORTED_EXTENSIONS = {
    ".txt", ".md", ".csv", ".json", ".docx",
    ".pdf", ".xlsx", ".xls", ".pptx", ".html", ".htm",
}
EMBEDDING_MODEL = "BAAI/bge-small-zh-v1.5"
RERANKER_MODEL = "BAAI/bge-reranker-base"

CHUNK_MAX_LEN = 500
CHUNK_OVERLAP = 100


def tokenize(text):
    return " ".join(w.strip() for w in jieba.cut(text) if w.strip())


class RAGEngine:
    """混合检索引擎：向量语义检索 + TF-IDF 关键词检索 + 增量同步。"""

    def __init__(self, data_dir="data", index_dir="index"):
        self.data_dir = Path(data_dir)
        self.data_dir.mkdir(parents=True, exist_ok=True)
        self.index_dir = Path(index_dir)
        self.index_dir.mkdir(parents=True, exist_ok=True)

        self._encoder = None
        self._reranker = None
        self._mode = "unknown"

        from retriever_store import NumpyTfidfStore, QdrantStore
        if os.environ.get("QDRANT_URL"):
            qd = QdrantStore(str(index_dir))
            self.store = qd if qd.available else NumpyTfidfStore(str(index_dir))
        else:
            self.store = NumpyTfidfStore(str(index_dir))

        self.sync = None

        self._load_or_sync()

    @property
    def encoder(self):
        if self._encoder is None:
            try:
                from sentence_transformers import SentenceTransformer
                print(f"[RAG] 加载嵌入模型 {EMBEDDING_MODEL} ...")
                self._encoder = SentenceTransformer(EMBEDDING_MODEL)
                self._mode = "vector"
                print("[RAG] 向量检索模式已启用")
            except Exception as e:
                print(f"[RAG] 嵌入模型加载失败: {e}")
                print("[RAG] 回退到 TF-IDF 模式")
                self._mode = "tfidf"
        return self._encoder

    @property
    def reranker(self):
        if self._reranker is None:
            try:
                from sentence_transformers import CrossEncoder
                print(f"[RAG] 加载重排序模型 {RERANKER_MODEL} ...")
                self._reranker = CrossEncoder(RERANKER_MODEL)
                print("[RAG] 重排序模型已启用")
            except Exception as e:
                print(f"[RAG] 重排序模型加载失败: {e}")
                print("[RAG] 跳过重排序，使用粗排结果")
        return self._reranker

    @property
    def chunks(self):
        return self.store.chunks

    @property
    def embeddings(self):
        return self.store.embeddings

    @property
    def vectorizer(self):
        return self.store.vectorizer

    @property
    def tfidf_matrix(self):
        return self.store.tfidf_matrix

    def _load_or_sync(self):
        loaded = self.store.load()
        if loaded:
            print(f"[RAG] 从磁盘加载索引: {self.store.count()} chunks")
        else:
            print("[RAG] 首次运行，执行全量同步...")

        from sync_pipeline import SyncPipeline
        self.sync = SyncPipeline(self, str(self.index_dir))

        report = self.sync.sync_all()
        total = report.get("total", {})
        if total.get("added") or total.get("updated") or total.get("deleted"):
            print(f"[RAG] 增量同步: +{total.get('added',0)} ~{total.get('updated',0)} -{total.get('deleted',0)}, 总计 {total.get('chunks',0)} chunks")
        elif not loaded:
            print(f"[RAG] 全量同步完成: {total.get('chunks',0)} chunks")

        if self._mode == "unknown":
            if self.store.embeddings is not None:
                self._mode = "vector"
            else:
                _ = self.encoder
                if self.store.embeddings is None and self.store.count() > 0:
                    self._embed_all_chunks()
                self.store.persist()

    def _embed_all_chunks(self):
        if self.encoder is None or self.store.count() == 0:
            return
        texts = [c["text"] for c in self.store.chunks]
        self.store.embeddings = self.encoder.encode(
            texts, show_progress_bar=True, normalize_embeddings=True
        )
        self.store.persist()

    def _add_chunks_to_store(self, chunks):
        """供 SyncPipeline 调用：将切块加入存储并 embed。"""
        if not chunks:
            return
        embeddings = None
        if self.encoder is not None:
            texts = [c["text"] for c in chunks]
            embeddings = self.encoder.encode(
                texts, show_progress_bar=False, normalize_embeddings=True
            )
        self.store.add_chunks(chunks, embeddings)

    def rebuild(self):
        """强制全量重同步。"""
        report = self.sync.full_resync()
        if self.store.embeddings is None and self.encoder is not None:
            self._embed_all_chunks()
        total = report.get("total", {})
        print(f"[RAG] 全量重建完成: {total.get('chunks',0)} chunks, 模式={self._mode}")

    def _split_chunks(self, text, source, meta, max_len=CHUNK_MAX_LEN, overlap=CHUNK_OVERLAP):
        paragraphs = re.split(r'\n\s*\n', text.strip())
        chunks = []
        prev_tail = ""

        for para in paragraphs:
            para = para.strip()
            if not para:
                continue

            if prev_tail:
                para = prev_tail + para
                prev_tail = ""

            if len(para) <= max_len:
                chunk_meta = dict(meta)
                chunks.append({"text": para, "source": source, "meta": chunk_meta})
                continue

            while len(para) > max_len:
                cut = para[:max_len]
                best_sep = -1
                for sep in ['。', '；', '！', '？', '\n', '.', '!', '?']:
                    pos = cut.rfind(sep)
                    if pos > max_len // 2:
                        best_sep = pos
                        break

                if best_sep > 0:
                    cut = para[:best_sep + 1]
                else:
                    best_space = cut.rfind(' ')
                    if best_space > max_len // 2:
                        cut = para[:best_space]

                chunk_meta = dict(meta)
                chunks.append({"text": cut, "source": source, "meta": chunk_meta})

                if overlap > 0 and len(cut) > overlap:
                    prev_tail = cut[-overlap:]
                    para = cut[-overlap:] + para[len(cut):]
                else:
                    para = para[len(cut):]

            if para:
                chunk_meta = dict(meta)
                chunks.append({"text": para, "source": source, "meta": chunk_meta})

        return chunks

    PRODUCT_KEYWORD_FILTER = {
        "耳机": ["耳机"],
        "体脂秤": ["体脂秤"],
        "电子秤": ["体脂秤"],
        "T恤": ["T恤", "短袖"],
        "牛仔裤": ["牛仔裤"],
        "跑鞋": ["跑鞋"],
        "书桌": ["书桌"],
        "面霜": ["面霜"],
        "冲锋衣": ["冲锋衣"],
        "降噪": ["降噪"],
    }

    def _keyword_filter(self, question, candidates):
        required = []
        for query_kw, doc_kws in self.PRODUCT_KEYWORD_FILTER.items():
            if query_kw in question:
                required.extend(doc_kws)

        if not required:
            return candidates

        filtered = [c for c in candidates
                    if any(kw in c["text"] for kw in required)]
        if filtered:
            return filtered
        return candidates

    @staticmethod
    def _expand_query(query):
        expansions = []
        height_map = {"一": "1", "二": "2", "三": "3", "四": "4", "五": "5",
                      "六": "6", "七": "7", "八": "8", "九": "9"}
        for m in re.finditer(r"一米(.)", query):
            digit = height_map.get(m.group(1))
            if digit:
                expansions.append(f"1{digit}0")

        term_map = {
            "男士": ["男装", "男"],
            "女士": ["女装", "女"],
            "多大码": ["尺码", "尺寸", "适合"],
            "穿什么": ["适合", "尺码"],
            "穿多大": ["适合", "尺码"],
            "多少钱": ["价格", "售价"],
            "包邮吗": ["运费", "包邮"],
            "能退吗": ["退换货", "退货"],
            "多久到": ["时效", "发货时间"],
            "有什么商品": ["商品", "产品"],
            "有哪些商品": ["商品", "产品"],
            "蓝牙耳机": ["耳机", "SC-505", "声科达", "降噪", "入耳式", "无线"],
            "耳机": ["耳机", "降噪", "入耳式"],
            "体脂秤": ["体脂秤", "体重", "BMI"],
            "电子秤": ["体脂秤", "体重", "BMI"],
            "T恤": ["T 恤", "短袖", "纯棉"],
            "裤子": ["男裤", "女裤", "休闲裤", "牛仔裤"],
            "裙子": ["连衣裙"],
            "上衣": ["T 恤", "卫衣", "衬衫"],
            "冲锋衣": ["童装", "防风", "儿童"],
            "跑鞋": ["跑鞋", "运动", "飞织", "减震"],
            "书桌": ["书桌", "实木", "办公"],
            "面霜": ["面霜", "保湿", "护肤"],
            "退换": ["退换货", "退货", "退款"],
            "换货": ["退换货", "换码"],
            "保修": ["质保", "售后"],
        }
        for colloquial, formal_list in term_map.items():
            if colloquial in query:
                expansions.extend(formal_list)

        if expansions:
            return query + " " + " ".join(set(expansions))
        return query

    def query(self, question, top_k=5, threshold=0.0):
        if self.store.count() == 0:
            return []

        expanded = self._expand_query(question)

        q_emb = None
        if self.encoder is not None and self.store.embeddings is not None:
            q_emb = self.encoder.encode(
                [expanded], normalize_embeddings=True
            )

        tokenized_q = tokenize(expanded) if self.store.vectorizer is not None else None

        candidates = self.store.search(
            query_embedding=q_emb,
            tokenized_query=tokenized_q,
            top_k=top_k * 3,
            threshold=threshold,
        )

        candidates = self._keyword_filter(question, candidates)

        if len(candidates) <= top_k:
            return candidates[:top_k]

        reranker = self.reranker
        if reranker is not None:
            try:
                pairs = [[question, c["text"]] for c in candidates]
                rerank_scores = reranker.predict(pairs)
                for i, c in enumerate(candidates):
                    c["rerank_score"] = float(rerank_scores[i])
                candidates.sort(key=lambda x: x["rerank_score"], reverse=True)
            except Exception as e:
                print(f"[RAG] 重排序失败，使用粗排结果: {e}")

        return candidates[:top_k]

    def list_documents(self):
        docs = []
        for filepath in sorted(self.data_dir.iterdir()):
            if filepath.is_file() and filepath.suffix.lower() in SUPPORTED_EXTENSIONS:
                stat = filepath.stat()
                docs.append({
                    "name": filepath.name,
                    "size": stat.st_size,
                    "modified": stat.st_mtime,
                })
        return docs

    def add_document(self, filename, content_bytes):
        from rag_engine import SUPPORTED_EXTENSIONS as SE
        filepath = self.data_dir / filename
        filepath.write_bytes(content_bytes)

        docs = self._load_file(filepath)
        new_chunks = []
        for text, meta in docs:
            new_chunks.extend(
                self._split_chunks(text, filepath.name, meta)
            )
        self._add_chunks_to_store(new_chunks)
        self.store.persist()
        return filepath.name

    def delete_document(self, filename):
        filepath = self.data_dir / filename
        if filepath.exists() and filepath.is_file():
            filepath.unlink()
            return self.store.delete_by_source(filename)
        return False

    def get_stats(self):
        return {
            "total_chunks": self.store.count(),
            "total_documents": len(self.list_documents()),
            "mode": self._mode if self._mode != "unknown" else "tfidf",
            "reranker_enabled": self._reranker is not None,
            "chunk_size": CHUNK_MAX_LEN,
            "chunk_overlap": CHUNK_OVERLAP,
            "index_persisted": (self.index_dir / "chunks.json").exists(),
            "supported_formats": sorted(SUPPORTED_EXTENSIONS),
            "store_type": type(self.store).__name__,
        }

    # ── 文件解析（仅用于初始导入 / 非 DB 管理的文档）──

    def _load_file(self, filepath):
        ext = filepath.suffix.lower()
        results = []

        if ext in (".txt", ".md"):
            text = filepath.read_text(encoding="utf-8")
            results.append((text, {"source": filepath.name, "type": "text"}))

        elif ext == ".csv":
            import csv
            with open(filepath, "r", encoding="utf-8", newline="") as f:
                reader = csv.DictReader(f)
                for i, row in enumerate(reader):
                    text = "；".join(f"{k}：{v}" for k, v in row.items() if v)
                    results.append((text, {
                        "source": filepath.name, "type": "csv", "row": i + 2
                    }))

        elif ext == ".json":
            data = json.loads(filepath.read_text(encoding="utf-8"))
            if isinstance(data, list):
                for i, item in enumerate(data):
                    text = self._json_to_text(item)
                    results.append((text, {
                        "source": filepath.name, "type": "json", "index": i
                    }))
            elif isinstance(data, dict):
                text = self._json_to_text(data)
                results.append((text, {
                    "source": filepath.name, "type": "json"
                }))

        elif ext == ".docx":
            from docx import Document
            doc = Document(str(filepath))
            parts = []
            for para in doc.paragraphs:
                t = para.text.strip()
                if t:
                    parts.append(t)
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append("；".join(cells))
            text = "\n\n".join(parts)
            results.append((text, {"source": filepath.name, "type": "docx"}))

        elif ext == ".pdf":
            try:
                import pdfplumber
                parts = []
                with pdfplumber.open(str(filepath)) as pdf:
                    for page in pdf.pages:
                        t = page.extract_text()
                        if t:
                            parts.append(t.strip())
                text = "\n\n".join(parts)
                results.append((text, {"source": filepath.name, "type": "pdf"}))
            except ImportError:
                try:
                    from PyPDF2 import PdfReader
                    reader = PdfReader(str(filepath))
                    parts = []
                    for page in reader.pages:
                        t = page.extract_text()
                        if t:
                            parts.append(t.strip())
                    text = "\n\n".join(parts)
                    results.append((text, {"source": filepath.name, "type": "pdf"}))
                except ImportError:
                    print(f"[RAG] PDF 解析库未安装，跳过 {filepath.name}")

        elif ext in (".xlsx", ".xls"):
            try:
                from openpyxl import load_workbook
                wb = load_workbook(str(filepath), data_only=True)
                parts = []
                for ws in wb.worksheets:
                    for row in ws.iter_rows(values_only=True):
                        cells = [str(c) for c in row if c is not None]
                        if cells:
                            parts.append("；".join(cells))
                text = "\n\n".join(parts)
                results.append((text, {"source": filepath.name, "type": "xlsx"}))
            except ImportError:
                print(f"[RAG] Excel 解析库未安装，跳过 {filepath.name}")

        elif ext == ".pptx":
            try:
                from pptx import Presentation
                prs = Presentation(str(filepath))
                parts = []
                for slide in prs.slides:
                    slide_texts = []
                    for shape in slide.shapes:
                        if hasattr(shape, "text") and shape.text.strip():
                            slide_texts.append(shape.text.strip())
                    if slide_texts:
                        parts.append("\n".join(slide_texts))
                text = "\n\n".join(parts)
                results.append((text, {"source": filepath.name, "type": "pptx"}))
            except ImportError:
                print(f"[RAG] PPTX 解析库未安装，跳过 {filepath.name}")

        elif ext in (".html", ".htm"):
            text = filepath.read_text(encoding="utf-8", errors="ignore")
            text = re.sub(r'<script[^>]*>.*?</script>', '', text, flags=re.DOTALL)
            text = re.sub(r'<style[^>]*>.*?</style>', '', text, flags=re.DOTALL)
            text = re.sub(r'<[^>]+>', ' ', text)
            text = re.sub(r'\s+', ' ', text).strip()
            results.append((text, {"source": filepath.name, "type": "html"}))

        return results

    @staticmethod
    def _json_to_text(obj, prefix=""):
        lines = []
        if isinstance(obj, dict):
            for k, v in obj.items():
                if isinstance(v, (dict, list)):
                    lines.append(f"{k}：{RAGEngine._json_to_text(v)}")
                else:
                    lines.append(f"{k}：{v}")
        elif isinstance(obj, list):
            for i, item in enumerate(obj):
                if isinstance(item, (dict, list)):
                    lines.append(f"[{i}] {RAGEngine._json_to_text(item)}")
                else:
                    lines.append(str(item))
        else:
            lines.append(str(obj))
        return "；".join(lines)
